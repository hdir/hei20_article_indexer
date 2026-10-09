"""Henter artiklene i startUrls.json og gjør dem om til LlamaIndex-Documents."""
import os
import re
import json
import logging
from typing import List
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field
from llama_index.core.schema import Document

from prompts import KEYWORD_PROMPT
from legacy_workflow import _invoke_with_retry

STOP_PHRASES = {
"vi har valgt ut dette for deg",
"spørsmål og svar som ligner",
"still oss et spørsmål",
"gå til spørsmål og svar",
"kontakt og hjelp",
"Fikk du svar på det du lurte på?" ,
"Les også:"
}

class KeywordResult(BaseModel):
    keywords: List[str] = Field(..., description="3-12 concise keywords/phrases in Norwegian")
    keyphrases: List[str] = Field(..., description="3-8 longer multiword phrases (2-6 words)")
    description: str = Field(..., description="1 sentence in norwegian")
    language: str = Field(..., description="ISO code like 'no'/'nb'/'nn'/'en'")


def _load_start_urls(json_path: str):
    """
    Reads a JSON file containing objects like:
       {"url": "...", "method": "GET"}
    and returns a list of URLs.
    """
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    # Extract urls, keep duplicates
    return [item["url"] for item in data if "url" in item]

def _pick_best_icon_url(soup, base_url: str) -> str | None:
    """
    Returns the best favicon for a page.
    Supports .png, .ico, .svg, .jpg, etc.
    """
    best_href = None
    best_score = -1

    for link in soup.find_all("link"):
        rel = link.get("rel") or []
        rel_str = " ".join(rel).lower() if isinstance(rel, list) else str(rel).lower()

        if "icon" not in rel_str:
            continue

        href = link.get("href")
        if not href:
            continue

        sizes = (link.get("sizes") or "").lower()
        m = re.search(r"(\d+)\s*x\s*(\d+)", sizes)

        # Scoring strategy
        if href.lower().endswith(".svg"):
            score = 10_000_000  # SVG is resolution-independent
        elif m:
            score = int(m.group(1)) * int(m.group(2))
        else:
            score = 16 * 16  # reasonable default

        if score > best_score:
            best_score = score
            best_href = href

    if best_href:
        return urljoin(base_url, best_href)

    # Apple touch icon (often PNG, high quality)
    apple = soup.find(
        "link",
        rel=lambda v: v and "apple-touch-icon" in " ".join(v).lower()
    )
    if apple and apple.get("href"):
        return urljoin(base_url, apple["href"])

    # Browser default fallback
    return urljoin(base_url, "/favicon.ico")

def _extract_braced_json(text: str, start_idx: int) -> str | None:
    depth = 0
    in_str = False
    esc = False
    for i in range(start_idx, len(text)):
        ch = text[i]
        if in_str:
            if esc:
                esc = False
            elif ch == "\\":
                esc = True
            elif ch == '"':
                in_str = False
        else:
            if ch == '"':
                in_str = True
            elif ch == "{":
                depth += 1
            elif ch == "}":
                depth -= 1
                if depth == 0:
                    return text[start_idx:i+1]
    return None

def _helsenorge_route_from_url(url: str) -> str:
    p = urlparse(url)
    frag = (p.fragment or "").strip()
    route = frag if frag.startswith("/") else (p.path or "/")
    if route != "/" and route.endswith("/"):
        route = route[:-1]
    return route

def _extract_helsenorge_selvtest_page(soup, url: str) -> dict | None:
    route = _helsenorge_route_from_url(url)

    script_text = ""
    for s in soup.find_all("script"):
        t = s.get_text() or ""
        if "window.selvtestConfig" in t:
            script_text = t
            break
    if not script_text:
        return None

    idx = script_text.find("window.selvtestConfig")
    idx = script_text.find("=", idx)
    idx = script_text.find("{", idx)
    if idx == -1:
        return None

    json_str = _extract_braced_json(script_text, idx)
    if not json_str:
        return None

    cfg = json.loads(json_str)
    pages = cfg.get("pages", [])

    # try exact route + common variants
    candidates = [route]
    if route != "/" and not route.endswith("/"):
        candidates.append(route + "/")
    if route.endswith("/") and route != "/":
        candidates.append(route[:-1])

    for cand in candidates:
        page = next((p for p in pages if p.get("urlPath") == cand), None)
        if page:
            return page
    return None

def _extract_page_data(url: str) -> dict:
    """Fetch a URL, parse HTML, extract text, title, and icon URL."""
    headers = {
        # Some sites behave better with a UA; harmless for ung.no too.
        "User-Agent": "Mozilla/5.0 (compatible; content-scraper/1.0)"
    }
    resp = requests.get(url, timeout=15, headers=headers)
    resp.raise_for_status()

    soup = BeautifulSoup(resp.text, "html.parser")

    title_tag = soup.find("title")
    title = title_tag.get_text(strip=True) if title_tag else ""

    description_tag = soup.find("meta", attrs={"name": "description"})
    description = description_tag.get("content") if description_tag else ""

    keywords_meta = soup.find("meta", attrs={"name": "keywords"})
    raw_keywords = keywords_meta.get("content", "") if keywords_meta else ""
    keywords = [kw.strip() for kw in raw_keywords.split(",") if kw.strip()] if raw_keywords else []

    # Main text content (default)
    body = soup.find("body")
    text = body.get_text(separator=" ", strip=True) if body else ""

    # Fallback for helsenorge.no where content lives in window.selvtestConfig
    host = (urlparse(url).hostname or "").lower()
    if "helsenorge.no" in host:
        page = _extract_helsenorge_selvtest_page(soup, url)
        if page:
            cfg_title = (page.get("title") or "").strip()
            if cfg_title:
                title = cfg_title

            parts = []
            for key in ("ingress1", "ingress2"):
                v = (page.get(key) or "").strip()
                if v:
                    parts.append(v)
            if parts:
                text = "\n\n".join(parts)

    return {
        "url": url,
        "title": title,
        "description": description,
        "content": text,
        "iconUrl": _pick_best_icon_url(soup, url),
        "keywords": keywords,
    }

def _split_text(text_to_split: str, split_sentences: set[str]) -> str:
    earliest = None
    for phrase in split_sentences:
        idx = text_to_split.find(phrase)
        if idx != -1 and (earliest is None or idx < earliest):
            earliest = idx
    return text_to_split if earliest is None else text_to_split[:earliest]

def _export_documents_to_excel(documents, out_xlsx_path: str) -> None:
    """
    One row per Document. Writes common metadata fields + doc_id + text length.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "documents"

    headers = [
        "doc_id",
        "url",
        "title",
        "description",
        "icon_url",
        "keywords",            # stored as comma-separated string
        "categories",
        "category",
        "severity",
        "answered_questions",  # stored as JSON string
        "text_length",
        "text"
    ]
    ws.append(headers)

    for doc in documents:
        md = getattr(doc, "metadata", {}) or {}
        keywords = md.get("keywords", [])
        if isinstance(keywords, list):
            keywords_str = ", ".join(str(k) for k in keywords)
        else:
            keywords_str = str(keywords or "")

        ws.append([
            getattr(doc, "doc_id", None) or md.get("url"),
            md.get("url"),
            md.get("title"),
            md.get("description"),
            md.get("icon_url"),
            keywords_str,
            json.dumps(md.get("categories", []), ensure_ascii=False),
            md.get("category", ""),
            md.get("severity", ""),
            json.dumps(md.get("answered_questions", []), ensure_ascii=False),
            len(getattr(doc, "text", "") or ""),
            getattr(doc, "text", ""),
        ])

    for col_idx, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(12, min(60, len(header) + 8))

    os.makedirs(os.path.dirname(out_xlsx_path) or ".", exist_ok=True)
    wb.save(out_xlsx_path)

def load_documents(llm, start_urls_file: str, output_dir: str) -> List[Document]:
    """
    Scrapes every URL in start_urls_file with requests + BeautifulSoup.
    Writes scraped_pages.json and documents.xlsx to output_dir for inspection.
    """
    start_urls = _load_start_urls(start_urls_file)
    logging.info("Scraping %d urls from %s", len(start_urls), start_urls_file)

    results = []
    for u in start_urls:
        try:
            results.append(_extract_page_data(u))
            print(f"Scraped: {u}")
        except Exception as e:
            print(f"Error scraping {u}: {e}")

    os.makedirs(output_dir, exist_ok=True)
    with open(os.path.join(output_dir, "scraped_pages.json"), "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)

    documents = []
    for item in results:
        # remove all text after the first stop phrase
        text_splitted = _split_text(item.get("content", ""), STOP_PHRASES)

        # if keywords are missing, use llm to build a set of keywords and update description as well
        keywords = item.get("keywords") or []
        description = item.get("description") or []
        if not keywords:
            prompt = KEYWORD_PROMPT.invoke({"context": text_splitted})
            kr: KeywordResult = _invoke_with_retry(llm.with_structured_output(KeywordResult), prompt)
            keywords = kr.keywords
            if not description:
                description = kr.description

        documents.append(Document(
            text=text_splitted,
            metadata={
                "url": item.get("url"),
                "title": item.get("title"),
                "description": description,
                "icon_url": item.get("iconUrl"),
                "keywords": keywords,
                "category": "",
                "severity": "",
                "answered_questions": []
            },
            doc_id=item.get("url"),
        ))

    _export_documents_to_excel(documents, out_xlsx_path=os.path.join(output_dir, "documents.xlsx"))
    return documents
